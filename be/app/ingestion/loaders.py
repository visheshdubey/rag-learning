import csv
import json
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxFile
from pptx import Presentation
from pypdf import PdfReader

from app.core.exceptions import LoaderError
from app.domain.models import Document


def _read_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="replace")


def _require_text(text: str, path: Path) -> str:
    cleaned = text.strip()
    if not cleaned:
        raise LoaderError(f"No extractable text in {path.name}")
    return cleaned


class TextLoader:
    def load(self, path: Path) -> list[Document]:
        text = _require_text(_read_text_file(path), path)
        return [
            Document(
                text=text,
                source=path.name,
                metadata={"file_type": path.suffix.lstrip(".").lower()},
            )
        ]


class PdfLoader:
    def load(self, path: Path) -> list[Document]:
        try:
            reader = PdfReader(str(path))
        except Exception as exc:
            raise LoaderError(f"Failed to read PDF {path.name}: {exc}") from exc

        documents: list[Document] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            documents.append(
                Document(
                    text=text,
                    source=path.name,
                    metadata={"file_type": "pdf", "page": index},
                )
            )
        if not documents:
            raise LoaderError(f"No extractable text in {path.name}")
        return documents


class DocxLoader:
    def load(self, path: Path) -> list[Document]:
        try:
            doc = DocxFile(str(path))
        except Exception as exc:
            raise LoaderError(f"Failed to read DOCX {path.name}: {exc}") from exc

        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        text = _require_text("\n".join(parts), path)
        return [
            Document(
                text=text,
                source=path.name,
                metadata={"file_type": "docx"},
            )
        ]


class PptxLoader:
    def load(self, path: Path) -> list[Document]:
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise LoaderError(f"Failed to read PPTX {path.name}: {exc}") from exc

        documents: list[Document] = []
        for index, slide in enumerate(presentation.slides, start=1):
            bits: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    bits.append(shape.text.strip())
            if not bits:
                continue
            documents.append(
                Document(
                    text="\n".join(bits),
                    source=path.name,
                    metadata={"file_type": "pptx", "slide": index},
                )
            )
        if not documents:
            raise LoaderError(f"No extractable text in {path.name}")
        return documents


class HtmlLoader:
    def load(self, path: Path) -> list[Document]:
        soup = BeautifulSoup(_read_text_file(path), "lxml")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        text = _require_text(soup.get_text(separator="\n"), path)
        return [
            Document(
                text=text,
                source=path.name,
                metadata={"file_type": "html"},
            )
        ]


class CsvLoader:
    def load(self, path: Path) -> list[Document]:
        try:
            with path.open(encoding="utf-8", newline="") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
        except Exception as exc:
            raise LoaderError(f"Failed to read CSV {path.name}: {exc}") from exc

        if not rows:
            raise LoaderError(f"No rows in {path.name}")

        documents: list[Document] = []
        for index, row in enumerate(rows, start=1):
            line = " | ".join(
                f"{key}: {value}" for key, value in row.items() if value
            )
            if not line:
                continue
            documents.append(
                Document(
                    text=line,
                    source=path.name,
                    metadata={"file_type": "csv", "row": index},
                )
            )
        if not documents:
            raise LoaderError(f"No extractable text in {path.name}")
        return documents


class JsonLoader:
    def load(self, path: Path) -> list[Document]:
        try:
            payload = json.loads(_read_text_file(path))
        except Exception as exc:
            raise LoaderError(f"Failed to read JSON {path.name}: {exc}") from exc

        if isinstance(payload, list) and payload and all(
            isinstance(item, dict) for item in payload
        ):
            documents = [
                Document(
                    text=json.dumps(item, ensure_ascii=False, indent=2),
                    source=path.name,
                    metadata={"file_type": "json", "index": index},
                )
                for index, item in enumerate(payload, start=1)
            ]
            return documents

        text = _require_text(json.dumps(payload, ensure_ascii=False, indent=2), path)
        return [
            Document(
                text=text,
                source=path.name,
                metadata={"file_type": "json"},
            )
        ]
